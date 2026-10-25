"""Build a results page (GitHub Actions job summary) from the scraped Excel files.

Usage: python summary.py <folder with .xlsx files> [download URL of the uploaded zip]
Prints Markdown; the workflow appends it to $GITHUB_STEP_SUMMARY.
"""
import glob
import os
import sys
from collections import Counter

from openpyxl import load_workbook

PLATFORM_BADGES = {'Swiggy': '🟠 Swiggy', 'Zomato': '🔴 Zomato'}


def rows_of(sheet):
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def block(rows, first_header):
    """Rows of the table whose header row starts with `first_header`, as dicts, up to the next blank row."""
    for index, row in enumerate(rows):
        if row and row[0] == first_header:
            headers = [h for h in row if h]
            records = []
            for data in rows[index + 1:]:
                if not any(value not in (None, '') for value in data):
                    break
                records.append(dict(zip(headers, data)))
            return records
    return []


def read_workbook(path):
    workbook = load_workbook(path, read_only=True)
    menu_rows = rows_of(workbook['Menu'])
    restaurant = block(menu_rows, 'Restaurant Name')[0]
    offers = [o for o in block(menu_rows, 'Offer') if o.get('Offer') != 'No offers found']
    items = block(menu_rows, 'Category')
    addons = block(rows_of(workbook['Add-ons']), 'Category') if 'Add-ons' in workbook.sheetnames else []
    timings = block(rows_of(workbook['Timings']), 'Day') if 'Timings' in workbook.sheetnames else []
    timings = [t for t in timings if t.get('Day') != 'No timings found']
    outlets = block(rows_of(workbook['Nearby Outlets']), 'Outlet Name') if 'Nearby Outlets' in workbook.sheetnames else []
    platform = os.path.basename(path).split('_', 1)[0]
    return {'file': os.path.basename(path), 'platform': platform, 'restaurant': restaurant, 'offers': offers,
            'items': items, 'addons': addons, 'timings': timings, 'outlets': outlets}


def cell(value):
    # Keep table cells on one line and stop '|' from breaking the table
    return ' '.join(str(value if value is not None else '').split()).replace('|', '/')


def table(headers, rows):
    lines = ['| ' + ' | '.join(headers) + ' |', '|' + '---|' * len(headers)]
    lines += ['| ' + ' | '.join(cell(value) for value in row) + ' |' for row in rows]
    return '\n'.join(lines)


def prices(items):
    return [item['Price'] for item in items if isinstance(item.get('Price'), (int, float)) and item['Price'] > 0]


def unique_items(items):
    # Swiggy repeats dishes in its Recommended section
    return list({(item.get('Item Name'), item.get('Price')): item for item in items}.values())


def hours_text(timings):
    by_day = {}
    for timing in timings:
        by_day.setdefault(timing['Day'], []).append(f"{timing.get('Opens', '')}–{timing.get('Closes', '')}")
    slots = {day: ', '.join(values) for day, values in by_day.items()}
    if len(set(slots.values())) == 1 and len(slots) == 7:
        return f'Every day: {next(iter(slots.values()))}'
    return ' · '.join(f'{day[:3]} {value}' for day, value in slots.items())


def restaurant_card(data):
    r, items = data['restaurant'], unique_items(data['items'])
    price_list = prices(items)
    lines = [f"## {PLATFORM_BADGES.get(data['platform'], data['platform'])} · {cell(r.get('Restaurant Name'))}", '']
    facts = [
        ('⭐ Rating', r.get('Restaurant Rating')),
        ('📍 Address', r.get('Restaurant Address')),
        ('🍽️ Cuisines', r.get('Cuisines')),
        ('💰 Cost for two', r.get('Cost for Two')),
        ('🛵 Delivery', r.get('Delivery Time')),
        ('🕒 Hours', hours_text(data['timings']) if data['timings'] else ''),
        ('🏷️ Offers', ' · '.join(filter(None, (
            ' '.join(filter(None, (o.get('Offer'), o.get('Coupon Code') and f"(code {o['Coupon Code']})")))
            for o in data['offers']))) or 'None listed'),
        ('🧾 FSSAI', r.get('FSSAI License No.')),
        ('🆔 Restaurant ID', r.get('Restaurant ID')),
    ]
    lines.append(table(['', ''], [(label, value) for label, value in facts if value]))
    lines.append('')
    if price_list:
        lines.append(f"**{len(items)} dishes** · prices ₹{min(price_list):g} – ₹{max(price_list):g} · "
                     f"average ₹{sum(price_list) / len(price_list):.0f} · **{len(data['addons'])} add-on options**")
        lines.append('')

    categories = Counter(item.get('Category') or 'Other' for item in items)
    lines += ['<details><summary>📂 Categories</summary>', '',
              table(['Category', 'Dishes', 'Price range'],
                    [(name, count, price_range([i for i in items if (i.get('Category') or 'Other') == name]))
                     for name, count in categories.most_common()]),
              '', '</details>', '']

    lines += ['<details><summary>📋 Full menu</summary>', '',
              table(['Category', 'Item', 'Price', 'Veg', 'Tags'],
                    [(i.get('Category'), i.get('Item Name'), f"₹{i['Price']:g}" if isinstance(i.get('Price'), (int, float)) else i.get('Price'),
                      '🟢' if str(i.get('Veg/Non-veg', '')).lower() == 'veg' else '🔴' if i.get('Veg/Non-veg') else '',
                      i.get('Tags')) for i in items]),
              '', '</details>', '']

    if data['outlets']:
        lines += [f"<details><summary>🏪 Other outlets ({len(data['outlets'])})</summary>", '',
                  table(['Locality', 'Delivery rating', 'Dining rating', 'Distance'],
                        [(o.get('Locality'), o.get('Delivery Rating'), o.get('Dining Rating'), o.get('Distance'))
                         for o in data['outlets']]),
                  '', '</details>', '']
    lines.append(f"<sub>📎 In the download: `{data['file']}`</sub>")
    return '\n'.join(lines)


def price_range(items):
    price_list = prices(items)
    if not price_list:
        return ''
    low, high = min(price_list), max(price_list)
    return f'₹{low:g}' if low == high else f'₹{low:g} – ₹{high:g}'


def comparison(results):
    rows = []
    for data in results:
        r, items = data['restaurant'], unique_items(data['items'])
        price_list = prices(items)
        rows.append((PLATFORM_BADGES.get(data['platform'], data['platform']), r.get('Restaurant Name'),
                     r.get('Restaurant Rating'), len(items),
                     f"₹{sum(price_list) / len(price_list):.0f}" if price_list else '',
                     ' · '.join(o.get('Offer', '') for o in data['offers']) or '—'))
    return table(['Platform', 'Restaurant', 'Rating', 'Dishes', 'Avg price', 'Offers'], rows)


def main():
    folder = sys.argv[1] if len(sys.argv) > 1 else '.'
    download_url = sys.argv[2] if len(sys.argv) > 2 else ''
    results = [read_workbook(path) for path in sorted(glob.glob(os.path.join(folder, '*.xlsx')))]

    out = ['# 🍽️ Menu Scraper', '']
    failed_file = os.path.join(folder, 'failed.txt')
    if os.path.exists(failed_file):
        with open(failed_file, encoding='utf-8') as f:
            failures = [line.rstrip('\n').split('\t', 1) for line in f if line.strip()]
        out += ['> [!WARNING]', f"> **{len(failures)} link{'s' if len(failures) != 1 else ''} not scraped**", '>']
        out += [f'> - {link} — {reason}' for link, reason in failures]
        out.append('')
    if not results:
        out.append('No menus were saved. Open the **scrape** job log below for details.')
        print('\n'.join(out))
        return
    scraped_on = results[0]['restaurant'].get('Scraped On', '')
    download = (f'### [⬇️ Download the Excel files (zip)]({download_url})' if download_url
                else '⬇️ Download the Excel files from **Artifacts** at the bottom of this page')
    out += [download, '', f"**{len(results)} menu{'s' if len(results) != 1 else ''} scraped** on {scraped_on}", '']
    if len(results) > 1:
        out += ['### 📊 Side by side', '', comparison(results), '']
    out += [restaurant_card(data) + '\n\n---\n' for data in results]
    print('\n'.join(out))


if __name__ == '__main__':
    main()
